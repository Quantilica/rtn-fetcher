"""Tests for Hive lake conversion of RTN Excel workbooks."""

from __future__ import annotations

from datetime import date
from importlib.metadata import entry_points
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Alignment

pytest.importorskip("quantilica.analytics.lake")
pytest.importorskip("duckdb")

from quantilica.analytics.lake import check_dataset, scan_dataset  # noqa: E402
from typer.testing import CliRunner  # noqa: E402

from rtn_fetcher.lake import (  # noqa: E402
    SOURCE_ID,
    converter,
    dataset_id_to_sheet,
    sheet_to_dataset_id,
)
from rtn_fetcher.plugin import app  # noqa: E402


def _build_synthetic_rtn_workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1.2-B"

    ws.append(["Voltar"])
    ws.append(["Tabela 1.2-B Resultado Primário do Governo Central - Mensal"])
    ws.append(["R$ Milhões - Valores de Mai/2026 - IPCA - Acumulado em 12 meses"])
    ws.append([])
    ws.append(["Discriminação", date(2020, 1, 1), date(2020, 2, 1), date(2020, 3, 1)])
    ws.append(["1. RECEITA TOTAL", 900000.0, 910000.0, 920000.0])
    ws["A6"].alignment = Alignment(indent=0)
    ws.append(["1.1 Receita Administrada", 500000.0, 505000.0, 510000.0])
    ws["A7"].alignment = Alignment(indent=1)
    ws.append(["Deflator - IPCA base Mai/2026", 5.54, 5.51, 5.48])
    ws["A8"].alignment = Alignment(indent=0)

    wb.save(path)
    return path


def test_entry_point_registered() -> None:
    eps = entry_points(group="quantilica.lake")
    matches = [ep for ep in eps if ep.name == "rtn"]
    assert len(matches) == 1
    loaded = matches[0].load()
    assert loaded is converter
    assert loaded.source_id == SOURCE_ID


def test_sheet_dataset_id_mapping() -> None:
    assert sheet_to_dataset_id("1.1") == "tabela-1-1"
    assert sheet_to_dataset_id("1.2-B") == "tabela-1-2-b"
    assert dataset_id_to_sheet("tabela-1-2-b") == "1.2-B"
    assert dataset_id_to_sheet("1.2-B") == "1.2-B"
    assert dataset_id_to_sheet("unknown") is None


def test_converter_datasets() -> None:
    datasets = converter.datasets()
    assert "tabela-1-1" in datasets
    assert "tabela-1-2-b" in datasets
    assert "tabela-2-1" in datasets


def test_convert_end_to_end(tmp_path: Path) -> None:
    raw_dir = tmp_path / "raw"
    raw_dir.mkdir()
    wb_path = raw_dir / "serie_historica_202610@20261021.xlsx"
    _build_synthetic_rtn_workbook(wb_path)

    lake_dir = tmp_path / "lake"
    res = converter.convert(raw_dir, lake_dir, datasets=["tabela-1-2-b"])
    assert not res.errors, res.errors
    assert res.converted == 1
    assert res.skipped == 0

    dataset_dir = lake_dir / SOURCE_ID / "tabela-1-2-b"
    assert check_dataset(dataset_dir) == []

    df = scan_dataset(dataset_dir).collect()
    assert df.height > 0
    assert "account" in df.columns
    assert "value" in df.columns

    # Second run skips
    res2 = converter.convert(raw_dir, lake_dir, datasets=["tabela-1-2-b"])
    assert not res2.errors
    assert res2.converted == 0
    assert res2.skipped == 1


def test_cli_convert(tmp_path: Path) -> None:
    raw_dir = tmp_path / "raw"
    raw_dir.mkdir()
    wb_path = raw_dir / "serie_historica_202610@20261021.xlsx"
    _build_synthetic_rtn_workbook(wb_path)

    lake_dir = tmp_path / "lake"
    runner = CliRunner()
    result = runner.invoke(
        app,
        [
            "convert",
            "-i",
            str(raw_dir),
            "--lake",
            str(lake_dir),
            "--dataset",
            "tabela-1-2-b",
        ],
    )
    assert result.exit_code == 0, result.output
    assert (lake_dir / SOURCE_ID / "tabela-1-2-b").exists()


def test_latest_workbook_prefers_newest_series(tmp_path: Path) -> None:
    from rtn_fetcher.lake import latest_workbook

    for rel in (
        "2015-03/Anexos RTN mar_2015.xlsx",
        "2026-07/serie_historica_jul26.xlsx",
        "2026-08/serie_historica_ago26.xlsx",
        "2026-08/~$serie_historica_ago26.xlsx",
        "2026-09/Anexo avulso.xlsx",
    ):
        (tmp_path / rel).parent.mkdir(parents=True, exist_ok=True)
        (tmp_path / rel).write_bytes(b"")
    assert [p.name for p in latest_workbook(tmp_path)] == ["serie_historica_ago26.xlsx"]
    assert latest_workbook(tmp_path / "vazio") == []
